"""Output-only, atomic exports of canonical reports and complete exception results."""

from __future__ import annotations

import csv
import os
import re
import tempfile
from collections.abc import Callable
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter

from auditor_support_tool.core.audit_procedure_models import ProcedureResult
from auditor_support_tool.core.audit_procedure_report_models import AuditProcedureReport
from auditor_support_tool.presentation.exception_export_table import (
    ExceptionExportTable,
    build_exception_export_table,
)

XLSX_MAX_ROWS = 1_048_576
XLSX_MAX_COLUMNS = 16_384
XLSX_MAX_TEXT = 32_767


class AuditExportError(Exception):
    """Controlled export failure without exposing evidence or settings."""


def default_export_filename(
    procedure_id: str, created_at: str, execution_id: str, kind: str, extension: str
) -> str:
    def safe(value: str) -> str:
        value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", value).strip(" .")[:60]
        if value.upper().split(".")[0] in {
            "CON",
            "PRN",
            "AUX",
            "NUL",
            *(f"{prefix}{i}" for prefix in ("COM", "LPT") for i in range(1, 10)),
        }:
            value = "_" + value
        return value or "Unknown"

    try:
        day = datetime.fromisoformat(created_at.replace("Z", "+00:00")).strftime("%Y%m%d")
    except ValueError:
        day = "Undated"
    return f"{safe(procedure_id)}_{safe(kind)}_{day}_{safe(execution_id[:8])}.{safe(extension)}"


def _csv_cell(value: object) -> object:
    # Harden only strings, not negative numeric values. Never mutate evidence.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _xlsx_cell(sheet, row: int, column: int, value: object) -> None:
    # Text avoids Excel's 15-digit precision limit for integers and decimals.
    if isinstance(value, (Decimal, float)) or (
        isinstance(value, int) and not isinstance(value, bool)
    ):
        value = str(value)
    if isinstance(value, str):
        if len(value.encode("utf-16-le")) // 2 > XLSX_MAX_TEXT:
            raise AuditExportError("A value exceeds Excel's cell text limit; use JSON or CSV.")
        if ILLEGAL_CHARACTERS_RE.search(value) or any(
            ord(char) in (0xFFFE, 0xFFFF) for char in value
        ):
            raise AuditExportError(
                "A value contains characters Excel cannot store; use JSON or CSV."
            )
    cell = sheet.cell(row, column)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # Explicit text, including formula-like strings and headers.


class AuditExportService:
    def _write(self, destination: Path, writer: Callable[[Path], None]) -> None:
        temporary = None
        try:
            destination = Path(destination)
            fd, name = tempfile.mkstemp(
                prefix=".audit-export-", suffix=destination.suffix, dir=destination.parent
            )
            temporary = Path(name)
            os.close(fd)
            writer(temporary)
            os.replace(temporary, destination)
        except AuditExportError:
            raise
        except PermissionError as error:
            raise AuditExportError(
                "Cannot write the export. The file may be open in Excel "
                "or the folder may be protected. Close it or choose another filename."
            ) from error
        except Exception as error:
            raise AuditExportError(
                "Export failed: the destination is unavailable or the data "
                "cannot be represented in this format. No records were skipped."
            ) from error
        finally:
            if temporary is not None:
                try:
                    temporary.unlink(missing_ok=True)
                except OSError as error:
                    raise AuditExportError(
                        "Export temporary-file cleanup failed. Check folder permissions."
                    ) from error

    def export_report_json(self, report: AuditProcedureReport, destination: Path) -> None:
        self._write(destination, lambda path: path.write_bytes(report.to_json().encode("utf-8")))

    def export_exceptions_csv(
        self, result: ProcedureResult | ExceptionExportTable, destination: Path
    ) -> None:
        def write(path: Path) -> None:
            table = (
                result
                if isinstance(result, ExceptionExportTable)
                else build_exception_export_table(result)
            )
            with path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.writer(stream)
                writer.writerow([_csv_cell(value) for value in table.headers])
                for row in table.rows:
                    writer.writerow([_csv_cell(value) for value in row])

        self._write(destination, write)

    def export_exceptions_xlsx(
        self, result: ProcedureResult | ExceptionExportTable, destination: Path
    ) -> None:
        def write(path: Path) -> None:
            table = (
                result
                if isinstance(result, ExceptionExportTable)
                else build_exception_export_table(result)
            )
            if len(table.rows) + 1 > XLSX_MAX_ROWS or len(table.headers) > XLSX_MAX_COLUMNS:
                raise AuditExportError(
                    "The export exceeds Excel's row or column limit; use CSV or JSON."
                )
            book = Workbook()
            try:
                sheet = book.active
                sheet.title = "Exceptions"
                for i, values in enumerate((table.headers, *table.rows), 1):
                    for j, value in enumerate(values, 1):
                        _xlsx_cell(sheet, i, j, value)
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = (
                    f"A1:{get_column_letter(len(table.headers))}{len(table.rows) + 1}"
                )
                for j, header in enumerate(table.headers, 1):
                    sheet.column_dimensions[get_column_letter(j)].width = min(
                        40, max(14, len(header) + 2)
                    )
                metadata = book.create_sheet("Metadata")
                for i, pair in enumerate((("Field", "Value"), *table.metadata), 1):
                    for j, value in enumerate(pair, 1):
                        _xlsx_cell(metadata, i, j, value)
                metadata.column_dimensions["A"].width = 24
                metadata.column_dimensions["B"].width = 80
                metadata.freeze_panes = "A2"
                book.save(path)
            finally:
                book.close()

        self._write(destination, write)
