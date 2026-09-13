"""PDF, DOCX and XLSX renderers of the same complete audit report document."""

from html import escape
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from auditor_support_tool.presentation.report_export_document import ReportDocument
from auditor_support_tool.services.audit_export_service import (
    XLSX_MAX_COLUMNS,
    XLSX_MAX_ROWS,
    AuditExportError,
    _xlsx_cell,
)


def _text(value: object) -> str:
    return "" if value is None else str(value)


def _html_table(headers, rows, title="") -> str:
    def cells(values, tag):
        return "".join(
            f"<{tag}>{escape(_text(v)).replace(chr(10), '<br/>')}</{tag}>" for v in values
        )

    return (
        '<table width="100%" cellspacing="0" cellpadding="5" border="1">'
        + "<thead>"
        + (
            f'<tr><th colspan="{len(headers)}" align="left" '
            f'style="font-size:14pt; background:#FFFFFF">{escape(title)}</th></tr>'
            if title
            else ""
        )
        + "<tr>"
        + cells(headers, "th")
        + "</tr></thead>"
        + "".join("<tr>" + cells(row, "td") + "</tr>" for row in rows)
        + "</table>"
    )


def _pdf_font_names() -> tuple[str, str]:
    """Use Windows Arial when available; fall back to PDF core fonts elsewhere."""
    windows_fonts = Path(r"C:\Windows\Fonts")
    regular = windows_fonts / "arial.ttf"
    bold = windows_fonts / "arialbd.ttf"
    if regular.exists() and bold.exists():
        try:
            if "ASTArial" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("ASTArial", str(regular)))
            if "ASTArial-Bold" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("ASTArial-Bold", str(bold)))
            return "ASTArial", "ASTArial-Bold"
        except Exception:
            pass
    return "Helvetica", "Helvetica-Bold"


def _pdf_markup(value: object) -> str:
    """Escape report evidence so source text is rendered literally, never as markup."""
    return escape(_text(value), quote=False).replace("\n", "<br/>")


def write_pdf(document: ReportDocument, path: Path) -> None:
    """Write a searchable, extractable, multi-page audit report PDF."""
    normal_font, bold_font = _pdf_font_names()
    page_width, _ = landscape(A4)
    usable_width = page_width - (32 * mm)
    styles = getSampleStyleSheet()
    normal = ParagraphStyle(
        "AuditNormal",
        parent=styles["BodyText"],
        fontName=normal_font,
        fontSize=8.5,
        leading=11,
        spaceAfter=3,
    )
    table_value = ParagraphStyle(
        "AuditTableValue",
        parent=normal,
        fontSize=8,
        leading=10,
        spaceAfter=0,
    )
    table_label = ParagraphStyle(
        "AuditTableLabel",
        parent=table_value,
        fontName=bold_font,
    )
    section_style = ParagraphStyle(
        "AuditSection",
        parent=styles["Heading1"],
        fontName=bold_font,
        fontSize=13,
        leading=16,
        spaceBefore=9,
        spaceAfter=5,
    )
    exception_style = ParagraphStyle(
        "AuditException",
        parent=styles["Heading2"],
        fontName=bold_font,
        fontSize=9.5,
        leading=12,
        spaceBefore=6,
        spaceAfter=3,
    )
    title_style = ParagraphStyle(
        "AuditTitle",
        parent=styles["Title"],
        fontName=bold_font,
        fontSize=20,
        leading=24,
        alignment=TA_CENTER,
        spaceAfter=12,
    )

    pdf = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title="Audit Procedure Report",
        author="Auditor Support Tool",
        creator="Auditor Support Tool",
    )
    story = [Paragraph("AUDIT PROCEDURE REPORT", title_style)]

    def field_table(rows) -> Table:
        values = [
            [
                Paragraph(_pdf_markup(label), table_label),
                Paragraph(_pdf_markup(value), table_value),
            ]
            for label, value in rows
        ]
        table = Table(values, colWidths=(0.28 * usable_width, 0.72 * usable_width))
        table.setStyle(
            TableStyle(
                [
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B7C3CC")),
                    ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EEF2F5")),
                ]
            )
        )
        return table

    def add_section(item) -> None:
        story.append(Paragraph(_pdf_markup(item.title), section_style))
        for paragraph in item.paragraphs:
            story.append(Paragraph(_pdf_markup(paragraph), normal))
        if item.rows:
            story.append(field_table(item.rows))
            story.append(Spacer(1, 4))

    for item in document.sections:
        add_section(item)

    story.append(Paragraph("Exception details", section_style))
    if not document.exception_rows:
        story.append(Paragraph("No exception records were identified.", normal))
    else:
        for group_number, (headers, rows) in enumerate(document.exception_bands(), 1):
            story.append(
                Paragraph(f"Exception fields - group {group_number}", exception_style)
            )
            header_cells = [Paragraph(_pdf_markup(value), table_label) for value in headers]
            body = [
                [Paragraph(_pdf_markup(value), table_value) for value in row]
                for row in rows
            ]
            remaining = len(headers) - 2
            shares = (0.08, 0.30, *((0.62 / remaining,) * remaining))
            table = Table(
                [header_cells, *body],
                colWidths=tuple(usable_width * share for share in shares),
                repeatRows=1,
                splitByRow=1,
            )
            table.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 3),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                        ("TOPPADDING", (0, 0), (-1, -1), 2),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#B7C3CC")),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6EC")),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 7))

    story.append(PageBreak())
    add_section(document.evidence)

    def footer(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFont(normal_font, 8)
        canvas.drawRightString(
            page_width - 16 * mm,
            8 * mm,
            f"Audit Procedure Report | Page {doc.page}",
        )
        canvas.restoreState()

    try:
        pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    except Exception as error:
        raise AuditExportError("PDF generation failed. Try Word or Excel.") from error

    if not path.exists() or path.stat().st_size <= 100:
        raise AuditExportError("PDF generation failed. Try Word or Excel.")
    with path.open("rb") as stream:
        if stream.read(5) != b"%PDF-":
            raise AuditExportError("PDF generation failed. Try Word or Excel.")


def write_docx(document: ReportDocument, path: Path) -> None:
    """Editable, macro-free Word document with repeated table headers."""
    book = Document()
    section = book.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.69)
    section.page_height = Inches(8.27)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = book.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = book.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for border in book.styles["Title"].element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
    book.add_heading("AUDIT PROCEDURE REPORT", 0)

    def table(headers, rows):
        grid = book.add_table(rows=1, cols=len(headers))
        grid.style = "Table Grid"
        grid.autofit = False
        width = section.page_width - section.left_margin - section.right_margin
        shares = (
            (0.32, 0.68)
            if len(headers) == 2
            else (0.08, 0.25, *((0.67 / (len(headers) - 2),) * (len(headers) - 2)))
        )
        for column, share in zip(grid.columns, shares, strict=True):
            column.width = int(width * share)
        for cell, value, share in zip(grid.rows[0].cells, headers, shares, strict=True):
            cell.width = int(width * share)
            cell.text = _text(value)
            for run in cell.paragraphs[0].runs:
                run.bold = True
        # python-docx has no public repeating-header API. Only this layout property is XML.
        grid.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for row in rows:
            for cell, value, share in zip(grid.add_row().cells, row, shares, strict=True):
                cell.width = int(width * share)
                cell.text = _text(value)
        book.add_paragraph()

    def content(item):
        book.add_heading(item.title, level=1)
        for paragraph in item.paragraphs:
            book.add_paragraph(paragraph)
        if item.rows:
            table(("Field", "Value"), item.rows)

    for item in document.sections:
        content(item)
    book.add_heading("Exception details", level=1)
    if not document.exception_rows:
        book.add_paragraph("No exception records were identified.")
    else:
        for index, (headers, rows) in enumerate(document.exception_bands(), 1):
            book.add_heading(f"Exception fields - group {index}", level=2)
            table(headers, rows)
    book.add_page_break()
    content(document.evidence)
    footer = section.footer.paragraphs[0]
    footer.text = "Audit Procedure Report | Page "
    field = OxmlElement("w:fldSimple")
    from docx.oxml.ns import qn

    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    book.save(path)


def write_xlsx(document: ReportDocument, path: Path) -> None:
    """Three analysis sheets; values always pass through the existing safe cell writer."""
    book = Workbook()
    book.remove(book.active)

    def sheet(name, headers, rows):
        if len(rows) + 1 > XLSX_MAX_ROWS or len(headers) > XLSX_MAX_COLUMNS:
            raise AuditExportError("The report exceeds Excel's row or column limit.")
        tab = book.create_sheet(name)
        for i, row in enumerate((headers, *rows), 1):
            for j, value in enumerate(row, 1):
                _xlsx_cell(tab, i, j, value)
                cell = tab.cell(i, j)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if i == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="24445C")
        tab.freeze_panes = "A2"
        if name == "Exceptions":
            tab.auto_filter.ref = tab.dimensions
        for j, header in enumerate(headers, 1):
            tab.column_dimensions[get_column_letter(j)].width = min(45, max(18, len(header) + 2))
        return tab

    def section_rows(sections):
        return tuple(
            row
            for item in sections
            for row in ((item.title, ""), *item.rows, *(("", p) for p in item.paragraphs), ("", ""))
        )

    try:
        summary = sheet(
            "Summary", ("AUDIT PROCEDURE REPORT", "Value"), section_rows(document.sections)
        )
        summary.column_dimensions["A"].width = 45
        summary.column_dimensions["B"].width = 100
        sheet("Exceptions", document.exception_headers, document.exception_rows)
        evidence = sheet(
            "Execution Evidence", ("Field", "Value"), section_rows((document.evidence,))
        )
        evidence.column_dimensions["A"].width = 36
        evidence.column_dimensions["B"].width = 100
        book.save(path)
    finally:
        book.close()
