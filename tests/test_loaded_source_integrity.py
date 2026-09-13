"""Loaded populations must retain the fingerprint of their parsed source bytes."""

from dataclasses import replace
from hashlib import sha256

import pytest

from auditor_support_tool.core.data_import_service import DataImportError, DataImportService
from auditor_support_tool.core.prepared_audit_dataset import PreparedAuditDataset
from auditor_support_tool.core.procedure_execution_status_service import (
    ProcedureExecutionStatus,
)
from auditor_support_tool.core.procedure_registry import ProcedureRegistry
from auditor_support_tool.core.source_integrity_service import SourceIntegrityService
from auditor_support_tool.core.test_engine_models import TestEngineStatus as EngineStatus
from auditor_support_tool.core.workbook_package import FieldMappingStatus
from auditor_support_tool.core.workbook_package_service import WorkbookPackageService
from auditor_support_tool.core.workspace_models import WorkspaceIdentity
from auditor_support_tool.core.workspace_service import WorkspaceService, WorkspaceServiceError
from auditor_support_tool.core.workspace_state import WorkspaceState
from auditor_support_tool.gui.pages.audit_procedures_page import AuditProceduresPage
from tests.test_source_integrity_service import create_paths, create_source_workbook
from tests.test_test_engine_service import StubProcedure, create_engine, create_source_file


def test_execution_requires_loaded_bytes_and_allows_controlled_reload(tmp_path):
    path = create_source_file(tmp_path)
    loader = WorkbookPackageService()
    source = PreparedAuditDataset(loader.build_package(path).datasets[0])
    original_hash = SourceIntegrityService().sha256_file(path)
    procedure = StubProcedure()
    engine = create_engine(procedure)

    def run(dataset):
        return engine.run(procedure_id="PROC001", source=dataset, source_path=path)

    unchanged = run(source)
    assert unchanged.status == EngineStatus.COMPLETED
    assert unchanged.result.context.source_sha256 == source.source_sha256 == original_hash

    path.write_text("reference,amount\nB001,200\nB002,300\n", encoding="utf-8")
    changed = run(source)
    assert changed.status == EngineStatus.FAILED
    assert not changed.was_executed
    assert changed.result is None
    assert procedure.run_count == 1
    assert "changed since it was loaded" in changed.error_message
    assert "Reload" in changed.error_message

    reloaded_source = PreparedAuditDataset(loader.build_package(path).datasets[0])
    reloaded = run(reloaded_source)
    assert reloaded.status == EngineStatus.COMPLETED
    assert reloaded.result.population_count == 2
    assert reloaded.result.context.source_sha256 == reloaded_source.source_sha256
    assert reloaded_source.source_sha256 != original_hash
    assert procedure.run_count == 2


def test_missing_loaded_fingerprint_fails_closed(tmp_path):
    path = create_source_file(tmp_path)
    dataset = WorkbookPackageService().build_package(path).datasets[0]
    dataset.loaded_table = replace(dataset.loaded_table, source_sha256="")
    procedure = StubProcedure()
    outcome = create_engine(procedure).run(
        procedure_id="PROC001", source=PreparedAuditDataset(dataset), source_path=path
    )
    assert outcome.status == EngineStatus.FAILED
    assert not outcome.was_executed
    assert procedure.run_count == 0
    assert "Reload" in outcome.error_message


@pytest.mark.parametrize("file_type", ("csv", "xlsx"))
def test_fingerprint_uses_the_bytes_parsed_even_if_path_changes(tmp_path, monkeypatch, file_type):
    path = create_source_file(tmp_path) if file_type == "csv" else create_source_workbook(tmp_path)
    original_bytes = path.read_bytes()
    loader = DataImportService()
    method = "_load_csv" if file_type == "csv" else "_load_excel"
    original_load = getattr(loader, method)

    def change_path_during_parse(*args, **kwargs):
        path.write_bytes(b"different source bytes")
        return original_load(*args, **kwargs)

    monkeypatch.setattr(loader, method, change_path_during_parse)
    table = loader.load_table(path)
    assert table.record_count == 1
    assert table.source_sha256 == sha256(original_bytes).hexdigest()
    assert table.source_sha256 != SourceIntegrityService().sha256_file(path)


def test_changed_source_does_not_create_execution_stamp(tmp_path, qtbot):
    path = create_source_file(tmp_path)
    package = WorkbookPackageService().build_package(path)
    package.datasets[0].mapping_status = FieldMappingStatus.CONFIRMED
    state = WorkspaceState()
    state.start_workspace(WorkspaceIdentity.create(name="Source binding"))
    state.set_workbook_package(package)
    procedure = StubProcedure()
    registry = ProcedureRegistry()
    registry.register(procedure)
    page = AuditProceduresPage(workspace_state=state, procedure_registry=registry)
    qtbot.addWidget(page)
    outcomes = []
    page.result_ready.connect(outcomes.append)
    path.write_text("reference,amount\nCHANGED,900\n", encoding="utf-8")
    with qtbot.waitSignal(page.result_ready, timeout=5000):
        page._run_procedure("PROC001")
    assert outcomes[-1].status == EngineStatus.FAILED
    assert procedure.run_count == 0
    assert state.procedure_execution_stamps == ()


def test_save_rejects_changed_source_before_replacing_managed_copy(tmp_path):
    path = create_source_file(tmp_path)
    state = WorkspaceState()
    state.start_workspace(WorkspaceIdentity.create(name="Source binding"))
    state.set_workbook_package(WorkbookPackageService().build_package(path))
    paths = create_paths(tmp_path)
    service = WorkspaceService(paths)
    saved = service.save_state(state, paths.workspaces / "audit.astworkspace")
    managed = paths.workspaces / "audit.astdata" / "source" / path.name
    previous_source, previous_document = managed.read_bytes(), saved.read_bytes()
    path.write_text("reference,amount\nCHANGED,900\n", encoding="utf-8")
    with pytest.raises(WorkspaceServiceError, match="Reload"):
        service.save_state(state)
    assert managed.read_bytes() == previous_source
    assert saved.read_bytes() == previous_document


def test_reopened_workspace_rebuilds_loaded_fingerprint(tmp_path):
    path = create_source_file(tmp_path)
    state = WorkspaceState()
    state.start_workspace(WorkspaceIdentity.create(name="Source binding"))
    state.set_workbook_package(WorkbookPackageService().build_package(path))
    service = WorkspaceService(create_paths(tmp_path))
    saved = service.save_state(state, tmp_path / "audit.astworkspace")
    reopened = WorkspaceState()
    service.load_into_state(reopened, saved)
    source = PreparedAuditDataset(reopened.active_dataset)
    outcome = create_engine(StubProcedure()).run(
        procedure_id="PROC001", source=source, source_path=reopened.source_path
    )
    assert outcome.status == EngineStatus.COMPLETED
    assert outcome.result.context.source_sha256 == source.source_sha256


def test_package_rejects_worksheets_loaded_from_different_source_versions(tmp_path, monkeypatch):
    from openpyxl import load_workbook

    path = create_source_workbook(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook.create_sheet("Reference")
    sheet.append(["Code"])
    sheet.append(["1000"])
    workbook.save(path)
    workbook.close()
    importer = DataImportService()
    original_load = importer.load_table
    calls = 0

    def change_between_worksheets(*args, **kwargs):
        nonlocal calls
        table = original_load(*args, **kwargs)
        calls += 1
        if calls == 1:
            # Appended bytes preserve a readable XLSX but change its fingerprint.
            path.write_bytes(path.read_bytes() + b"changed between worksheets")
        return table

    monkeypatch.setattr(importer, "load_table", change_between_worksheets)
    with pytest.raises(DataImportError, match="changed while loading"):
        WorkbookPackageService(import_service=importer).build_package(path)


def test_failed_rerun_does_not_report_previous_success_as_completed(
    tmp_path,
    qtbot,
):
    """A failed latest attempt must not appear Completed because an older success exists."""

    path = create_source_file(tmp_path)

    package = WorkbookPackageService().build_package(path)
    dataset = package.datasets[0]
    dataset.mapping_status = FieldMappingStatus.CONFIRMED

    state = WorkspaceState()
    state.start_workspace(
        WorkspaceIdentity.create(
            name="Latest execution status",
        )
    )
    state.set_workbook_package(package)

    procedure = StubProcedure()

    registry = ProcedureRegistry()
    registry.register(procedure)

    page = AuditProceduresPage(
        workspace_state=state,
        procedure_registry=registry,
    )
    qtbot.addWidget(page)

    outcomes = []
    page.result_ready.connect(outcomes.append)

    # First execution succeeds and creates valid historical evidence.
    with qtbot.waitSignal(page.result_ready, timeout=5000):
        page._run_procedure("PROC001")

    assert outcomes[-1].status == EngineStatus.COMPLETED

    successful_stamp = state.get_procedure_execution_stamp(
        "PROC001",
        dataset.dataset_id,
    )

    assert successful_stamp is not None

    # The same procedure is attempted again, but this attempt fails.
    procedure.behavior = "raise"

    with qtbot.waitSignal(page.result_ready, timeout=5000):
        page._run_procedure("PROC001")

    assert outcomes[-1].status == EngineStatus.FAILED

    # Preserve the previous successful run as reproducibility evidence.
    assert (
        state.get_procedure_execution_stamp(
            "PROC001",
            dataset.dataset_id,
        )
        == successful_stamp
    )

    # But the procedure list must represent the latest failed attempt
    # as requiring another run rather than showing Completed.
    status = page._procedure_execution_status(
        definition=procedure.definition,
        source=PreparedAuditDataset(dataset),
    )

    assert status == ProcedureExecutionStatus.NEEDS_RERUN
