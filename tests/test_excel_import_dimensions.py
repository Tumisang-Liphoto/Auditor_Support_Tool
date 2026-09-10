"""Regression coverage for read-only Excel worksheets with unreliable bounds."""

from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from auditor_support_tool.core.data_import_service import DataImportError, DataImportService
from auditor_support_tool.core.data_models import SOURCE_ROW_FIELD
from auditor_support_tool.core.workbook_package_service import WorkbookPackageService


def make_workbook(tmp_path: Path, rows: list[list], dimensions: str, suffix: str = ".xlsx") -> Path:
    """Use write-only output to omit dimensions, optionally supplying a false range."""
    workbook = Workbook(write_only=dimensions != "normal")
    worksheet = workbook.active if dimensions == "normal" else workbook.create_sheet()
    worksheet.title = "General_Ledger"
    for row in rows:
        worksheet.append(row)
    path = tmp_path / f"{dimensions}{suffix}"
    workbook.save(path)
    if dimensions == "understated":
        with ZipFile(path) as archive:
            entries = [(info, archive.read(info.filename)) for info in archive.infolist()]
        with ZipFile(path, "w") as archive:
            for info, data in entries:
                if info.filename == "xl/worksheets/sheet1.xml":
                    data = data.replace(b"<sheetData>", b'<dimension ref="A1:A1"/><sheetData>')
                archive.writestr(info, data)
    return path


@pytest.mark.parametrize("dimensions", ["normal", "missing", "understated"])
@pytest.mark.parametrize("suffix", [".xlsx", ".xlsm"])
def test_import_preserves_headers_values_and_source_rows(tmp_path, dimensions, suffix):
    path = make_workbook(
        tmp_path,
        [
            ["Invoice", "", "Invoice", SOURCE_ROW_FIELD],
            ["INV-001", "Context", "INV-001", "original", 100],
            [],
            ["INV-002", "Other", "INV-002", "source", 250],
        ],
        dimensions,
        suffix,
    )
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if dimensions == "missing":
            assert workbook.active.max_row is None
            assert workbook.active.max_column is None
        elif dimensions == "understated":
            assert workbook.active.calculate_dimension() == "A1:A1"
    finally:
        workbook.close()
    service = DataImportService()
    info = service.inspect_source(path).worksheets[0]
    assert (info.maximum_row, info.maximum_column, info.estimated_data_rows) == (4, 5, 3)
    table = service.load_table(path, worksheet_name="General_Ledger")
    assert table.file_type == suffix[1:]
    assert table.headers == (
        "Invoice",
        "Unnamed Column 2",
        "Invoice [2]",
        f"{SOURCE_ROW_FIELD} [source]",
        "Unnamed Column 5",
    )
    assert table.original_headers == ("Invoice", "", "Invoice", SOURCE_ROW_FIELD, "")
    assert len(table.summary.header_changes) == 4
    assert table.record_count == 2
    assert table.summary.source_records_read == 3
    assert table.summary.blank_rows_skipped == 1
    assert [row[SOURCE_ROW_FIELD] for row in table.rows] == [2, 4]
    assert [row["Unnamed Column 5"] for row in table.rows] == [100, 250]
    assert table.rows[0][f"{SOURCE_ROW_FIELD} [source]"] == "original"
    package = WorkbookPackageService().build_package(path)
    dataset = package.get_dataset_by_worksheet("General_Ledger")
    assert dataset is not None
    assert dataset.data_profile.record_count == 2
    assert [column.source_column for column in dataset.columns] == list(table.headers)
    assert len({column.column_id for column in dataset.columns}) == 5


@pytest.mark.parametrize("dimensions", ["normal", "missing", "understated"])
@pytest.mark.parametrize("rows", [[], [["Invoice"]], [["Invoice", "Amount"]]])
def test_empty_and_header_only_sheets(tmp_path, dimensions, rows):
    path = make_workbook(tmp_path, rows, dimensions)
    service = DataImportService()
    assert service.inspect_source(path).worksheets[0].estimated_data_rows == 0
    assert WorkbookPackageService().build_package(path).datasets == []
    if rows:
        table = service.load_table(path)
        assert table.headers == tuple(rows[0])
        assert table.record_count == 0
    else:
        with pytest.raises(DataImportError, match="worksheet is empty"):
            service.load_table(path)
    with pytest.raises(DataImportError, match="header row is beyond"):
        service.load_table(path, header_row=2)


@pytest.mark.parametrize("dimensions", ["normal", "missing", "understated"])
def test_offset_header_and_cached_formula_values(tmp_path, dimensions):
    path = make_workbook(
        tmp_path, [[], ["Invoice", "Amount"], ["INV-001", "=1+2"], [], ["INV-002", 50]], dimensions
    )
    table = DataImportService().load_table(path, header_row=2)
    assert table.headers == ("Invoice", "Amount")
    assert [row[SOURCE_ROW_FIELD] for row in table.rows] == [3, 5]
    # openpyxl writes no formula cache: data_only must not expose or evaluate it.
    assert [row["Amount"] for row in table.rows] == [None, 50]
    assert table.summary.blank_rows_skipped == 1
