"""Complete auditor report formats, source fidelity, safety and format selection."""

from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
from zipfile import ZipFile

import pytest
from docx import Document
from openpyxl import load_workbook
from PySide6.QtPdf import QPdfDocument
from PySide6.QtWidgets import QRadioButton

from auditor_support_tool.gui.dialogs.report_export_dialog import ReportExportDialog
from auditor_support_tool.presentation.exception_source_records import ExceptionSourceRecords
from auditor_support_tool.presentation.report_export_document import (
    ReportExportRequest,
    build_report_document,
)
from auditor_support_tool.services.audit_export_service import AuditExportError, AuditExportService
from tests.test_audit_procedure_report import _definition, _result


@pytest.fixture
def report_request():
    result = _result()
    sources = ExceptionSourceRecords(
        ("Original Ref", "Value", "Narrative"),
        {
            (e.source_record_id, e.source_row_number): (f" 00{i} ", Decimal("12.340"), "=1+1")
            for i, e in enumerate(result.exception_records)
        },
        result.context.dataset_id,
        result.context.source_sha256,
    )
    return ReportExportRequest(
        _definition(), result, sources, "Annual review", "Example auditee", "Ledger", "Transactions"
    )


def pdf_text(path):
    pdf = QPdfDocument()
    assert pdf.load(str(path)) == QPdfDocument.Error.None_
    assert pdf.pageCount() > 0
    text = "\n".join(pdf.getAllText(i).text() for i in range(pdf.pageCount()))
    pdf.close()
    return text


def docx_text(path):
    book = Document(path)
    return "\n".join(
        [p.text for p in book.paragraphs]
        + [cell.text for table in book.tables for row in table.rows for cell in row.cells]
    )


def test_format_selector_exact_options_and_pdf_default(qtbot):
    dialog = ReportExportDialog()
    qtbot.addWidget(dialog)
    assert [r.text() for r in dialog.findChildren(QRadioButton)] == [
        "PDF",
        "Word (.docx)",
        "Excel (.xlsx)",
    ]
    assert dialog.format == "pdf"
    dialog._formats["docx"].click()
    assert dialog.format == "docx"


@pytest.mark.parametrize("format", ("pdf", "docx", "xlsx"))
def test_complete_formats_preserve_values_and_evidence(report_request, tmp_path, qapp, format):
    before = deepcopy(report_request)
    model = build_report_document(report_request)
    path = tmp_path / ("report." + format)
    AuditExportService().export_report(report_request, path, format)
    assert path.stat().st_size > 100
    if format == "pdf":
        assert path.read_bytes().startswith(b"%PDF-")
        text = pdf_text(path)
    elif format == "docx":
        with ZipFile(path) as archive:
            assert archive.testzip() is None
            assert not any("vba" in name.lower() for name in archive.namelist())
        text = docx_text(path)
    else:
        book = load_workbook(path)
        assert book.sheetnames == ["Summary", "Exceptions", "Execution Evidence"]
        assert book["Exceptions"].max_row == len(report_request.result.exception_records) + 1
        assert book["Exceptions"].freeze_panes == "A2"
        assert book["Exceptions"].auto_filter.ref
        rows = list(book["Exceptions"].values)
        assert rows[1][2:5] == (" 000 ", "12.340", "=1+1")
        assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
        summary = dict(book["Summary"].values)
        assert summary["Exceptions"] == str(report_request.result.exception_count)
        text = "\n".join(
            str(c.value) for sheet in book for row in sheet for c in row if c.value is not None
        )
        book.close()
    for value in (
        "Annual review",
        "Example auditee",
        "Original Ref",
        "12.340",
        "=1+1",
        report_request.result.context.source_sha256,
        report_request.result.context.mapping_fingerprint,
        model.report.report_fingerprint,
        *(e.source_record_id for e in report_request.result.exception_records),
    ):
        assert value in text
    assert report_request == before
    assert build_report_document(report_request).report.to_json() == model.report.to_json()


@pytest.mark.parametrize("fault", ("missing", "hash", "dataset", "row", "invalid", "width"))
def test_incomplete_or_invalid_evidence_leaves_destination_untouched(
    report_request, tmp_path, fault
):
    sources = report_request.sources
    if fault == "missing":
        sources = None
    elif fault == "hash":
        sources = replace(sources, source_sha256="f" * 64)
    elif fault == "dataset":
        sources = replace(sources, dataset_id="another")
    else:
        records = dict(sources.records)
        key = next(iter(records))
        if fault == "row":
            records.pop(key)
        elif fault == "invalid":
            records[key] = ("x", float("nan"), "text")
        else:
            records[key] = ("x",)
        sources = replace(sources, records=records)
    path = tmp_path / "report.xlsx"
    path.write_bytes(b"existing report")
    with pytest.raises(AuditExportError, match="evidence"):
        AuditExportService().export_report(replace(report_request, sources=sources), path, "xlsx")
    assert path.read_bytes() == b"existing report"
    assert list(tmp_path.iterdir()) == [path]


@pytest.mark.parametrize("format,extension", (("json", "json"), ("docx", "doc"), ("pdf", "xlsx")))
def test_unsupported_report_format_or_extension(report_request, tmp_path, format, extension):
    with pytest.raises(AuditExportError, match="matching"):
        AuditExportService().export_report(
            report_request, tmp_path / ("report." + extension), format
        )
    assert not list(tmp_path.iterdir())


@pytest.mark.parametrize("format", ("pdf", "docx", "xlsx"))
def test_writer_failure_is_atomic(report_request, tmp_path, monkeypatch, format):
    from auditor_support_tool.services import audit_report_writers

    def fail(document, path):
        path.write_bytes(b"partial")
        raise RuntimeError("confidential server detail")

    monkeypatch.setattr(audit_report_writers, "write_" + format, fail)
    path = tmp_path / ("report." + format)
    path.write_bytes(b"original")
    with pytest.raises(AuditExportError) as error:
        AuditExportService().export_report(report_request, path, format)
    assert "confidential" not in str(error.value)
    assert path.read_bytes() == b"original"
    assert list(tmp_path.iterdir()) == [path]


@pytest.mark.parametrize("format", ("pdf", "docx", "xlsx"))
def test_multipage_full_population_and_wide_fields(report_request, tmp_path, qapp, format):
    original = report_request.result.exception_records[0]
    exceptions = tuple(
        replace(
            original,
            source_row_number=i + 2,
            source_record_id=f"dataset-gl:row-{i + 2}",
            reason=f"Exception marker {i:03}",
        )
        for i in range(65)
    )
    result = replace(report_request.result, exception_records=exceptions, exception_count=65)
    headers = tuple(f"Source field {i}" for i in range(9))
    sources = replace(
        report_request.sources,
        headers=headers,
        records={
            (e.source_record_id, e.source_row_number): tuple(
                f"raw {j} record {i:03}" for j in range(9)
            )
            for i, e in enumerate(exceptions)
        },
    )
    path = tmp_path / ("wide." + format)
    AuditExportService().export_report(
        replace(report_request, result=result, sources=sources), path, format
    )
    if format == "pdf":
        text = pdf_text(path)
    elif format == "docx":
        text = docx_text(path)
    else:
        book = load_workbook(path)
        text = "\n".join(str(c.value) for row in book["Exceptions"] for c in row)
        assert book["Exceptions"].max_row == 66
        book.close()
    for i in range(65):
        assert f"Exception marker {i:03}" in text
        assert f"raw 8 record {i:03}" in text


@pytest.mark.parametrize("format", ("pdf", "docx", "xlsx"))
def test_zero_exceptions_and_literal_markup(report_request, tmp_path, qapp, format):
    request = replace(
        report_request,
        result=replace(report_request.result, exception_records=(), exception_count=0),
        sources=None,
        workspace_name="<b>Literal & workspace</b>",
    )
    path = tmp_path / ("zero." + format)
    AuditExportService().export_report(request, path, format)
    if format == "pdf":
        assert "<b>Literal & workspace</b>" in pdf_text(path)
    elif format == "docx":
        assert "<b>Literal & workspace</b>" in docx_text(path)
    else:
        book = load_workbook(path)
        assert book["Exceptions"].max_row == 1
        assert dict(book["Summary"].values)["Audit / workspace"] == "<b>Literal & workspace</b>"
        book.close()


@pytest.mark.parametrize("format", ("pdf", "docx", "xlsx"))
def test_control_characters_fail_without_silent_loss(report_request, tmp_path, format):
    request = replace(report_request, workspace_name="invalid\x00label")
    path = tmp_path / ("bad." + format)
    with pytest.raises(AuditExportError, match="invalid"):
        AuditExportService().export_report(request, path, format)
    assert not path.exists()
