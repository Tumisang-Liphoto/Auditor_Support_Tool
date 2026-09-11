"""File exports preserve canonical evidence and existing destinations."""

import csv
import json
from copy import deepcopy
from dataclasses import replace

import pytest
from openpyxl import load_workbook

from auditor_support_tool.core.audit_procedure_report_builder import AuditProcedureReportBuilder
from auditor_support_tool.services import audit_export_service as exports
from tests.test_audit_procedure_report import _definition, _result
from tests.test_exception_export_table import export_result


def test_canonical_json_repeat_and_provenance(tmp_path):
    report = AuditProcedureReportBuilder().build(definition=_definition(), result=_result())
    path = tmp_path / "report.json"
    service = exports.AuditExportService()
    for _ in range(2):
        service.export_report_json(report, path)
        assert path.read_bytes() == report.to_json().encode("utf-8")
    payload = json.loads(path.read_bytes())
    assert payload == report.to_dict()
    assert len(payload["exceptions"]) == 2
    assert payload["exceptions"][0]["source_row_number"] == 12
    assert payload["report_fingerprint"] == report.report_fingerprint
    assert not {"settings", "credentials", "ai", "generated_at"} & payload.keys()


def test_csv_complete_fidelity_and_safety(tmp_path):
    result = export_result()
    for i, value in enumerate(["=1", " +1", "\t-1", "\n@x", "plain"]):
        result.exception_records[0].values[f"text{i}"] = value
    original = deepcopy(result)
    path = tmp_path / "rows.csv"
    exports.AuditExportService().export_exceptions_csv(result, path)
    assert path.read_bytes().startswith(b"\xef\xbb\xbf")
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 2
    first = rows[0]
    assert first["source_record_id"] == result.exception_records[0].source_record_id
    assert first["value.formula"] == "'  =SUM(A1:A2)"
    assert first["value.negative_number"] == "-12.30"
    assert first["value.negative_string"] == "'-12.30"
    assert first["value.source_record_id"] == "0000123"
    assert first["value.quoted"] == 'é, "quoted"\nline'
    assert first["value.null"] == ""
    assert json.loads(first["value.nested"]) == {"a": True, "z": [1, "é"]}
    for i in range(4):
        assert first[f"value.text{i}"].startswith("'")
    assert first["value.text4"] == "plain"
    assert result == original


def test_xlsx_complete_text_and_metadata(tmp_path):
    result = export_result()
    original = deepcopy(result)
    path = tmp_path / "rows.xlsx"
    exports.AuditExportService().export_exceptions_xlsx(result, path)
    book = load_workbook(path)
    try:
        assert book.sheetnames == ["Exceptions", "Metadata"]
        sheet = book["Exceptions"]
        assert sheet.max_row == 3
        headers = [cell.value for cell in sheet[1]]
        cells = dict(zip(headers, sheet[2], strict=True))
        assert cells["value.formula"].value == "  =SUM(A1:A2)"
        assert cells["value.formula"].data_type == "s"
        assert cells["value.source_record_id"].value == "0000123"
        assert cells["value.negative_number"].value == "-12.30"
        assert cells["source_row_number"].value == "12"
        assert cells["value.nested"].value == '{"a":true,"z":[1,"é"]}'
        assert sheet.freeze_panes == "A2"
        assert sheet.auto_filter.ref.endswith("3")
        metadata = dict(book["Metadata"].values)
        assert metadata["exception_count"] == "2"
        assert metadata["audit_period_start"] == result.context.audit_period_start
        assert metadata["source_sha256"] == result.context.source_sha256
    finally:
        book.close()
    assert result == original


@pytest.mark.parametrize("format", ["csv", "xlsx"])
def test_zero_exports(tmp_path, format):
    result = replace(_result(), exception_records=(), exception_count=0, exception_rate=0)
    path = tmp_path / ("empty." + format)
    getattr(exports.AuditExportService(), "export_exceptions_" + format)(result, path)
    if format == "csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            assert len(list(csv.reader(stream))) == 1
    else:
        book = load_workbook(path)
        assert book["Exceptions"].max_row == 1
        assert dict(book["Metadata"].values)["exception_count"] == "0"
        book.close()


@pytest.mark.parametrize("failure", ["replace", "serialize", "write"])
def test_atomic_failure_preserves_destination(tmp_path, monkeypatch, failure):
    path = tmp_path / "report.json"
    path.write_bytes(b"original")
    report = AuditProcedureReportBuilder().build(definition=_definition(), result=_result())
    if failure == "replace":

        def fail(*args):
            raise PermissionError("secret response")

        monkeypatch.setattr(exports.os, "replace", fail)
    elif failure == "serialize":
        report.metrics["bad"] = object()
    else:

        def fail(*args):
            raise OSError("secret response")

        monkeypatch.setattr(exports.Path, "write_bytes", fail)
    with pytest.raises(exports.AuditExportError) as error:
        exports.AuditExportService().export_report_json(report, path)
    assert "secret response" not in str(error.value)
    assert path.read_bytes() == b"original"
    assert list(tmp_path.iterdir()) == [path]


@pytest.mark.parametrize(
    "value",
    ["x" * 32768, "bad\x00text", "=1", 10**20],
    ids=["overlong", "illegal-character", "formula", "large-integer"],
)
def test_xlsx_limits_and_literal_formula(tmp_path, value):
    result = _result()
    result.exception_records[0].values["check"] = value
    path = tmp_path / "data.xlsx"
    if isinstance(value, str) and (len(value) > 32767 or "\x00" in value):
        path.write_bytes(b"old")
        with pytest.raises(exports.AuditExportError):
            exports.AuditExportService().export_exceptions_xlsx(result, path)
        assert path.read_bytes() == b"old"
        assert list(tmp_path.iterdir()) == [path]
    else:
        exports.AuditExportService().export_exceptions_xlsx(result, path)
        book = load_workbook(path)
        sheet = book["Exceptions"]
        index = [c.value for c in sheet[1]].index("value.check") + 1
        assert sheet.cell(2, index).value == str(value)
        assert sheet.cell(2, index).data_type == "s"
        book.close()


@pytest.mark.parametrize("limit", ["XLSX_MAX_ROWS", "XLSX_MAX_COLUMNS"])
def test_xlsx_dimension_limits(tmp_path, monkeypatch, limit):
    monkeypatch.setattr(exports, limit, 1)
    with pytest.raises(exports.AuditExportError, match="limit"):
        exports.AuditExportService().export_exceptions_xlsx(_result(), tmp_path / "limit.xlsx")
    assert list(tmp_path.iterdir()) == []


def test_default_name_safe_stable():
    name = exports.default_export_filename(
        "CON", "2026-09-11T12:00:00Z", "a1b2c3d4-more", "AllExceptions", "csv"
    )
    assert name == "_CON_AllExceptions_20260911_a1b2c3d4.csv"
    name = exports.default_export_filename(
        'x<>:"/\\|?*.', "invalid", "bad/id", "AllExceptions", "xlsx"
    )
    assert not any(char in name for char in '<>:"/\\|?*')
